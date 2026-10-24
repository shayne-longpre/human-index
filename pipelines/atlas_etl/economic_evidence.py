"""Derive source-native economic tables without temporal or geographic splicing."""

from __future__ import annotations

import argparse
import csv
import json
import math
from collections import Counter
from pathlib import Path

import openpyxl

from .acquire_economic_sources import WORLD_BANK_INDICATORS, sha256, world_bank_response


TRANSFORM_ID = "economic-source-native-table-v1"
TRANSFORM_VERSION = "1.0.0"
MADDISON_ID = "maddison-gdppc-2011-international-dollars"
WB_SOURCE_ID = "world-bank-wdi-economics-2026-10-08"
MADDISON_SOURCE_ID = "maddison-project-database-2023-v1"


def verified_artifacts(lockfile: Path) -> dict[str, dict]:
    lock = json.loads(lockfile.read_text())
    artifacts = {item["id"]: item for item in lock["artifacts"]}
    if len(artifacts) != len(lock["artifacts"]):
        raise ValueError("Duplicate economic source artifact ID")
    for item in artifacts.values():
        path = Path(item["path"])
        if not path.exists() or sha256(path) != item["sha256"] or path.stat().st_size != item["bytes"]:
            raise ValueError(f"Economic source artifact changed or missing: {item['id']}")
    return artifacts


def _value(value: object, *, allow_negative: bool = False) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"Invalid economic source value: {value!r}")
    if not allow_negative and value < 0:
        raise ValueError(f"Negative value in a nonnegative economic indicator: {value}")
    return float(value)


def _source_sheet(workbook: openpyxl.Workbook, name: str, header_row: int) -> list[dict]:
    sheet = workbook[name]
    references = []
    active_code = None
    for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
        code, period, citation = (list(row[:3]) + [None] * 3)[:3]
        if isinstance(code, str) and len(code) == 3:
            active_code = code
        if not isinstance(citation, str) or not citation.strip():
            continue
        references.append({
            "source_sheet": name,
            "source_row": row_number,
            "countrycode": active_code,
            "period_as_written": period,
            "citation_as_written": citation,
        })
    return references


def maddison_dataset(artifact: dict) -> tuple[dict, list[list[object]]]:
    workbook = openpyxl.load_workbook(artifact["path"], read_only=True, data_only=True)
    try:
        expected = {"Notes", "Sources", "GDPpc", "Population", "Full data", "Maddison original sources"}
        if not expected.issubset(workbook.sheetnames):
            raise ValueError("Maddison 2023 workbook sheets changed")
        notes = [value for row in workbook["Notes"].iter_rows(values_only=True) for value in row]
        if "Real GDP per capita in 2011$" not in notes or "Population, mid-year (thousands)" not in notes:
            raise ValueError("Maddison source units differ from audited workbook")
        sheet = workbook["Full data"]
        iterator = sheet.iter_rows(values_only=True)
        if tuple(next(iterator)[:6]) != ("countrycode", "country", "region", "year", "gdppc", "pop"):
            raise ValueError("Maddison full-data columns changed")
        supports: dict[str, dict] = {}
        seen: set[tuple[str, int]] = set()
        full_rows = []
        known = []
        counts: Counter[int] = Counter()
        for source_row, row in enumerate(iterator, 2):
            code, country, region, year, raw_gdp, raw_population = row[:6]
            if not isinstance(code, str) or len(code) != 3 or not isinstance(country, str):
                raise ValueError(f"Invalid Maddison source support at Full data!A{source_row}")
            if isinstance(year, bool) or not isinstance(year, int) or year < 1 or year > 2022:
                raise ValueError(f"Invalid Maddison year at Full data!D{source_row}")
            key = (code, year)
            if key in seen:
                raise ValueError(f"Duplicate Maddison country/year: {key}")
            seen.add(key)
            support = {"code": code, "name": country, "region": region}
            if code in supports and supports[code] != support:
                raise ValueError(f"Maddison source support label changed for {code}")
            supports[code] = support
            value = _value(raw_gdp)
            population_thousands = _value(raw_population)
            full_rows.append([code, year, value, population_thousands, source_row])
            if value is not None:
                known.append([code, year, value, source_row, population_thousands])
                counts[year] += 1
        references = _source_sheet(workbook, "Sources", 4) + _source_sheet(workbook, "Maddison original sources", 3)
        if not known or not references:
            raise ValueError("Maddison workbook has no publishable values or original-source references")
        result = {
            "schema_version": "0.1.0",
            "id": MADDISON_ID,
            "source_release_id": MADDISON_SOURCE_ID,
            "source_artifact_id": artifact["id"],
            "source_artifact_sha256": artifact["sha256"],
            "source_sheet": "Full data",
            "indicator": "gdppc",
            "label": "GDP per capita, Maddison 2023",
            "unit": "2011 international dollars per person",
            "support_type": "Maddison source country/economy code; not a historical-polity border",
            "temporal_policy": "exact_source_year_only",
            "transformation": {"id": TRANSFORM_ID, "version": TRANSFORM_VERSION, "policy": "unaltered_nonnull_source_values; no_splice_or_interpolation"},
            "source_note": "The workbook calls this Real GDP per capita in 2011$. Country-labelled historical estimates are not individual incomes. Source-sheet references are retained as country-level candidates, not asserted as exact per-row provenance.",
            "source_citation_url": "https://www.rug.nl/ggdc/historicaldevelopment/maddison/releases/maddison-project-database-2023?lang=en",
            "source_supports": sorted(supports.values(), key=lambda item: item["code"]),
            "source_references": references,
            "global_source_note": "GDP pc: 2008 - 2022: Total Economy Database the Conference Board for all countries included in TED. Otherwise UN national accounts statistics",
            "source_row_count": len(full_rows),
            "known_value_count": len(known),
            "available_years": sorted(counts),
            "known_by_year": {str(year): counts[year] for year in sorted(counts)},
            "observations": known,
        }
        return result, full_rows
    finally:
        workbook.close()


def world_bank_dataset(artifact: dict, countries_artifact: dict, code: str, source_update: str) -> tuple[dict, list[list[object]]]:
    meta, rows = world_bank_response(Path(artifact["path"]), code)
    if meta.get("lastupdated") != source_update:
        raise ValueError(f"Unexpected World Bank source update for {code}")
    _, country_rows = world_bank_response(Path(countries_artifact["path"]), None)
    supports = {
        row["id"]: {"code": row["id"], "name": row["name"], "region": row["region"]["value"]}
        for row in country_rows if row.get("region", {}).get("id") != "NA"
    }
    if len(supports) != 217:
        raise ValueError("World Bank nonaggregate support universe changed; review required")
    seen: set[tuple[str, int]] = set()
    full_rows = []
    known = []
    counts: Counter[int] = Counter()
    for source_row, row in enumerate(rows, 1):
        source_code = row.get("countryiso3code")
        if source_code not in supports:
            continue
        year = int(row["date"])
        key = (source_code, year)
        if key in seen:
            raise ValueError(f"Duplicate World Bank {code} country/year: {key}")
        seen.add(key)
        value = _value(row.get("value"), allow_negative=code in {"FP.CPI.TOTL.ZG", "NE.GDI.TOTL.ZS"})
        full_rows.append([source_code, year, value, source_row])
        if value is not None:
            known.append([source_code, year, value, source_row])
            counts[year] += 1
    label, unit = WORLD_BANK_INDICATORS[code]
    if not known:
        raise ValueError(f"No known World Bank values for {code}")
    result = {
        "schema_version": "0.1.0",
        "id": f"wdi-{code.lower().replace('.', '-')}",
        "source_release_id": WB_SOURCE_ID,
        "source_artifact_id": artifact["id"],
        "source_artifact_sha256": artifact["sha256"],
        "source_sheet": "World Bank API response data array",
        "indicator": code,
        "label": label,
        "unit": unit,
        "support_type": "World Bank nonaggregate country/territory code; separate from historical-polity geography",
        "temporal_policy": "exact_source_year_only",
        "transformation": {"id": TRANSFORM_ID, "version": TRANSFORM_VERSION, "policy": "unaltered_nonnull_source_values; exclude_API_aggregates; no_splice_or_interpolation"},
        "source_note": "This is the World Bank's country/territory indicator at its source support. It is not a local or person-level measurement.",
        "source_citation_url": f"https://data.worldbank.org/indicator/{code}",
        "source_supports": sorted(supports.values(), key=lambda item: item["code"]),
        "source_references": [],
        "source_row_count": len(full_rows),
        "known_value_count": len(known),
        "available_years": sorted(counts),
        "known_by_year": {str(year): counts[year] for year in sorted(counts)},
        "observations": known,
    }
    return result, full_rows


def derive(lockfile: Path, output: Path) -> list[dict]:
    artifacts = verified_artifacts(lockfile)
    lock = json.loads(lockfile.read_text())
    if lock["world_bank_source_update"] != "2026-10-08":
        raise ValueError("World Bank release date changed; review indicator definitions before publication")
    output.mkdir(parents=True, exist_ok=True)
    catalog = []
    countries = artifacts["world_bank_economic_country_metadata"]
    sources = [(maddison_dataset(artifacts["maddison_2023_workbook"]), "maddison_2023_workbook")]
    sources += [
        (world_bank_dataset(artifacts[f"world_bank_economic_{code.lower().replace('.', '_')}"] , countries, code, lock["world_bank_source_update"]), f"world_bank_economic_{code.lower().replace('.', '_')}")
        for code in WORLD_BANK_INDICATORS
    ]
    for (dataset, full_rows), artifact_id in sources:
        stem = dataset["id"]
        json_path = output / f"{stem}.json"
        csv_path = output / f"{stem}.csv"
        json_path.write_text(json.dumps(dataset, separators=(",", ":"), ensure_ascii=False) + "\n", encoding="utf-8")
        with csv_path.open("w", newline="", encoding="utf-8") as stream:
            writer = csv.writer(stream, lineterminator="\n")
            if artifact_id == "maddison_2023_workbook":
                writer.writerow(["source_release_id", "indicator", "source_support_code", "year", "gdppc_2011_international_dollars", "source_population_thousands", "value_status", "original_record_locator"])
                for code, year, value, population, source_row in full_rows:
                    writer.writerow([dataset["source_release_id"], "gdppc", code, year, value if value is not None else "", population if population is not None else "", "known" if value is not None else "unknown", f"Full data!A{source_row}:F{source_row}"])
            else:
                writer.writerow(["source_release_id", "indicator", "source_support_code", "year", "value", "unit", "value_status", "original_record_locator"])
                for code, year, value, source_row in full_rows:
                    writer.writerow([dataset["source_release_id"], dataset["indicator"], code, year, value if value is not None else "", dataset["unit"], "known" if value is not None else "unknown", f"API data[{source_row - 1}]"])
        catalog.append({
            "id": dataset["id"], "source_release_id": dataset["source_release_id"],
            "source_artifact_id": artifact_id, "source_artifact_sha256": dataset["source_artifact_sha256"],
            "indicator": dataset["indicator"], "label": dataset["label"], "unit": dataset["unit"],
            "support_type": dataset["support_type"], "temporal_policy": dataset["temporal_policy"],
            "available_years": dataset["available_years"], "source_row_count": dataset["source_row_count"],
            "known_value_count": dataset["known_value_count"],
            "data_file": f"economics/{json_path.name}", "download_file": f"economics/{csv_path.name}",
        })
    return catalog


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--lockfile", type=Path, default=Path("data/registry/economic-source-lock.json"))
    parser.add_argument("--output", type=Path, default=Path("data/derived/economics"))
    args = parser.parse_args()
    result = derive(args.lockfile, args.output)
    print(json.dumps({"datasets": len(result), "known_observations": sum(item["known_value_count"] for item in result)}))


if __name__ == "__main__":
    main()
