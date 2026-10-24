"""Derive source-native PWT 11 tables; no ratio or cross-source rebasing."""

from __future__ import annotations

import argparse
import csv
import json
import math
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from .acquire_economic_sources import sha256
from .economic_evidence import TRANSFORM_VERSION


SOURCE_ID = "penn-world-table-11.0"
TRANSFORM_ID = "pwt-source-native-table-v1"
FIELDS = {
    "rgdpe": ("Expenditure-side real GDP, PWT 11", "million 2021 US$ at chained PPPs"),
    "rgdpo": ("Output-side real GDP, PWT 11", "million 2021 US$ at chained PPPs"),
    "emp": ("Persons engaged, PWT 11", "million persons"),
    "avh": ("Annual hours per person engaged, PWT 11", "hours per person engaged"),
    "ctfp": ("Total factor productivity level, PWT 11", "USA = 1 at current PPPs"),
    "cn": ("Capital stock, PWT 11", "million 2021 US$ at current PPPs"),
    "labsh": ("Labor compensation share of GDP, PWT 11", "fraction of GDP"),
}
EXPECTED_LEGEND = {
    "rgdpe": "Expenditure-side real GDP at chained PPPs (in mil. 2021US$)",
    "rgdpo": "Output-side real GDP at chained PPPs (in mil. 2021US$)",
    "emp": "Number of persons engaged (in millions)",
    "avh": "Average annual hours worked by persons engaged",
    "ctfp": "TFP level at current PPPs (USA=1)",
    "cn": "Capital stock at current PPPs (in mil. 2021US$)",
    "labsh": "Share of labour compensation in GDP at current national prices",
}


def derive(lockfile: Path, output: Path) -> list[dict]:
    lock = json.loads(lockfile.read_text())
    raw = Path(lock["path"])
    if (lock["id"] != "penn-world-table-11.0-workbook" or sha256(raw) != lock["sha256"]
            or raw.stat().st_size != lock["bytes"]):
        raise ValueError("Pinned PWT workbook changed or missing")
    workbook = openpyxl.load_workbook(raw, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["Info", "Legend", "Data"] or workbook["Info"]["A1"].value != "Penn World Table, version 11.0":
            raise ValueError("PWT release workbook structure changed")
        legend = {row[0]: row[1] for row in workbook["Legend"].iter_rows(max_row=60, values_only=True) if row[0] in FIELDS}
        if legend != EXPECTED_LEGEND:
            raise ValueError("PWT variable definitions changed")
        sheet = workbook["Data"]
        header = tuple(next(sheet.iter_rows(max_row=1, values_only=True)))
        if header[:4] != ("countrycode", "country", "currency_unit", "year") or not set(FIELDS).issubset(header):
            raise ValueError("PWT data columns changed")
        columns = {field: header.index(field) for field in FIELDS}
        flag_columns = {flag: header.index(flag) for flag in ("i_cig", "i_outlier")}
        # The XLSX declares 1,048,576 rows although the reviewed release has exactly 185*74 records.
        raw_rows = list(sheet.iter_rows(min_row=2, max_row=13691, values_only=True))
        if len(raw_rows) != 13690 or any(value is not None for value in next(sheet.iter_rows(min_row=13692, max_row=13692, values_only=True))):
            raise ValueError("PWT source row extent changed")
        supports = {}
        seen = set()
        support_rows: Counter[str] = Counter()
        for index, row in enumerate(raw_rows, 2):
            code, name, _, year = row[:4]
            if not isinstance(code, str) or len(code) != 3 or not isinstance(name, str) or not isinstance(year, int) or not 1950 <= year <= 2023:
                raise ValueError(f"Invalid PWT support/year at Data!{index}")
            if code in supports and supports[code] != name:
                raise ValueError(f"PWT country label changed within {code}")
            supports[code] = name
            support_rows[code] += 1
            if (code, year) in seen:
                raise ValueError(f"Duplicate PWT support/year {code} {year}")
            seen.add((code, year))
        if len(supports) != 185 or any(count != 74 for count in support_rows.values()):
            raise ValueError("PWT support universe changed")
        output.mkdir(parents=True, exist_ok=True)
        catalog = []
        for field, (label, unit) in FIELDS.items():
            column = columns[field]
            letter = get_column_letter(column + 1)
            observations = []
            counts: Counter[int] = Counter()
            full_rows = []
            for index, row in enumerate(raw_rows, 2):
                value = row[column]
                if value is not None:
                    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0:
                        raise ValueError(f"Invalid PWT {field} at Data!{letter}{index}")
                    if field == "labsh" and value > 1:
                        raise ValueError(f"Invalid PWT labor share at Data!{letter}{index}")
                    observations.append([row[0], row[3], float(value), index])
                    counts[row[3]] += 1
                full_rows.append((row, value, index))
            dataset_id = f"pwt11-{field}"
            dataset = {
                "schema_version": "0.1.0", "id": dataset_id,
                "source_release_id": SOURCE_ID, "source_artifact_id": "pwt_11_workbook",
                "source_artifact_sha256": lock["sha256"], "source_sheet": "Data", "source_column": letter,
                "indicator": field, "label": label, "unit": unit,
                "support_type": "PWT source country code; not a historical-polity border",
                "temporal_policy": "exact_source_year_only",
                "transformation": {"id": TRANSFORM_ID, "version": TRANSFORM_VERSION,
                                   "policy": "unaltered_nonnull_source_values; no_ratios_or_geographic_splice"},
                "source_note": f"PWT Legend defines {field}: {legend[field]}. The source i_cig and i_outlier flags are retained in the CSV; other quality fields remain in the pinned raw workbook. This is a country-level economic context, not an individual distribution.",
                "source_citation_url": "https://www.rug.nl/ggdc/productivity/pwt/?lang=en",
                "source_supports": [{"code": code, "name": name, "region": None} for code, name in sorted(supports.items())],
                "source_references": [],
                "global_source_note": "Cite Feenstra, Inklaar and Timmer (2015) as required by the PWT publisher.",
                "source_row_count": len(raw_rows), "known_value_count": len(observations),
                "available_years": sorted(counts),
                "known_by_year": {str(year): counts[year] for year in sorted(counts)},
                "observations": observations,
            }
            json_path = output / f"{dataset_id}.json"
            csv_path = output / f"{dataset_id}.csv"
            json_path.write_text(json.dumps(dataset, separators=(",", ":"), ensure_ascii=False) + "\n")
            with csv_path.open("w", newline="", encoding="utf-8") as stream:
                writer = csv.writer(stream, lineterminator="\n")
                writer.writerow(["source_release_id", "indicator", "source_support_code", "year", "value", "unit", "value_status", "i_cig", "i_outlier", "original_record_locator"])
                for row, value, index in full_rows:
                    writer.writerow([SOURCE_ID, field, row[0], row[3], value if value is not None else "", unit,
                                     "known" if value is not None else "unknown", row[flag_columns["i_cig"]],
                                     row[flag_columns["i_outlier"]], f"Data!{letter}{index}"])
            catalog.append({
                "id": dataset_id, "source_release_id": SOURCE_ID,
                "source_artifact_id": "pwt_11_workbook", "source_artifact_sha256": lock["sha256"],
                "indicator": field, "label": label, "unit": unit,
                "support_type": dataset["support_type"], "temporal_policy": "exact_source_year_only",
                "available_years": dataset["available_years"], "source_row_count": len(raw_rows),
                "known_value_count": len(observations), "data_file": f"economics/{json_path.name}",
                "download_file": f"economics/{csv_path.name}",
            })
        return catalog
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--lockfile", type=Path, default=Path("data/registry/pwt-source-lock.json"))
    parser.add_argument("--output", type=Path, default=Path("data/derived/pwt"))
    args = parser.parse_args()
    result = derive(args.lockfile, args.output)
    print(json.dumps({"datasets": len(result), "known_values": sum(item["known_value_count"] for item in result)}))


if __name__ == "__main__":
    main()
