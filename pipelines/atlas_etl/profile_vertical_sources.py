"""Profile newly acquired vertical sources without promoting them to a release.

The output is a review artifact: it records source-native supports and obvious
quality gates so later adapters cannot treat every workbook or CSV as the same
kind of country-year evidence.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import tempfile
import warnings
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .core import sha256, write_json
from .sources.world_bank_economics import WorldBankEconomicsAdapter


def _as_int(value: object) -> int:
    return int(float(str(value)))


def _clio_profile(path: Path) -> dict[str, Any]:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Data Long Format"]
    # Clio's generated files incorrectly declare A1 as the used range even
    # though rows and columns are present. Resetting dimensions lets openpyxl
    # stream the actual source cells without editing the workbook.
    sheet.reset_dimensions()
    rows = sheet.values
    headers = tuple(next(rows))
    if headers[:4] != ("ccode", "country.name", "year", "value"):
        raise RuntimeError(f"Unexpected Clio long-format columns in {path}")
    row_count = 0
    known_values = 0
    codes: set[int] = set()
    years: set[int] = set()
    duplicate_keys = 0
    keys: set[tuple[int, int]] = set()
    for row in rows:
        if not row or row[0] is None or row[2] is None:
            continue
        code = _as_int(row[0])
        year = _as_int(row[2])
        row_count += 1
        known_values += row[3] is not None
        codes.add(code)
        years.add(year)
        key = (code, year)
        duplicate_keys += key in keys
        keys.add(key)

    metadata_sheet = workbook["Metadata"]
    metadata_sheet.reset_dimensions()
    metadata = {
        str(row[0]): str(row[1])
        for row in metadata_sheet.values
        if row and len(row) > 1 and row[0] is not None and row[1] is not None
    }
    workbook.close()
    return {
        "filename": path.name,
        "sha256": sha256(path),
        "rows": row_count,
        "known_values": known_values,
        "country_codes": len(codes),
        "year_range": [min(years), max(years)] if years else None,
        "duplicate_country_year_keys": duplicate_keys,
        "source_citation": metadata.get("Text Citation"),
        "source_download_url": metadata.get("Downloaded from"),
        "publication_status": "review_only_terms_and_2012_boundary_support_unresolved",
    }


def _csv_rows(archive_path: Path, member: str) -> Iterable[dict[str, str]]:
    archive = zipfile.ZipFile(archive_path)
    source = archive.open(member)
    text = io.TextIOWrapper(source, encoding="utf-8-sig", newline="")
    try:
        yield from csv.DictReader(text)
    finally:
        text.close()
        archive.close()


def _ucdp_ged_profile(path: Path) -> dict[str, Any]:
    years: set[int] = set()
    country_ids: set[str] = set()
    regions: Counter[str] = Counter()
    violence_types: Counter[str] = Counter()
    rows = 0
    nonmonotonic_fatality_fields = 0
    for row in _csv_rows(path, "GEDEvent_v26_1.csv"):
        rows += 1
        years.add(int(row["year"]))
        country_ids.add(row["country_id"])
        regions[row["region"]] += 1
        violence_types[row["type_of_violence"]] += 1
        low, best, high = int(row["low"]), int(row["best"]), int(row["high"])
        nonmonotonic_fatality_fields += not low <= best <= high
    return {
        "filename": path.name,
        "sha256": sha256(path),
        "event_rows": rows,
        "year_range": [min(years), max(years)],
        "country_ids": len(country_ids),
        "rows_by_region": dict(sorted(regions.items())),
        "rows_by_type_of_violence": dict(sorted(violence_types.items())),
        "rows_not_low_le_best_le_high": nonmonotonic_fatality_fields,
        "source_support": "georeferenced_event_location",
        "publication_status": "adapter_pending_geography_and_measure_decision",
    }


def _ucdp_acd_profile(path: Path) -> dict[str, Any]:
    years: set[int] = set()
    conflicts: set[str] = set()
    locations: set[str] = set()
    intensity: Counter[str] = Counter()
    rows = 0
    for row in _csv_rows(path, "UcdpPrioConflict_v26_1.csv"):
        rows += 1
        years.add(int(row["year"]))
        conflicts.add(row["conflict_id"])
        locations.add(row["location"])
        intensity[row["intensity_level"]] += 1
    return {
        "filename": path.name,
        "sha256": sha256(path),
        "conflict_year_rows": rows,
        "conflicts": len(conflicts),
        "location_labels": len(locations),
        "year_range": [min(years), max(years)],
        "rows_by_intensity_level": dict(sorted(intensity.items())),
        "source_support": "conflict_location_year_and_parties",
        "publication_status": "adapter_pending_geography_and_measure_decision",
    }


def _seshat_profile(path: Path) -> dict[str, Any]:
    with tempfile.TemporaryDirectory() as directory, zipfile.ZipFile(path) as archive:
        workbook_member = next(name for name in archive.namelist() if name.endswith(".xlsx"))
        archive.extract(workbook_member, directory)
        workbook_path = Path(directory) / workbook_member
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)

        sheet_profiles: dict[str, dict[str, int]] = {}
        for sheet in workbook.worksheets:
            populated_rows = 0
            max_populated_columns = 0
            for row in sheet.values:
                populated = [index for index, value in enumerate(row, start=1) if value is not None]
                if populated:
                    populated_rows += 1
                    max_populated_columns = max(max_populated_columns, populated[-1])
            sheet_profiles[sheet.title] = {
                "populated_rows_including_header": populated_rows,
                "max_populated_columns": max_populated_columns,
            }

        time_sheet = workbook["TSDat123"]
        time_rows = time_sheet.values
        header = tuple(next(time_rows))
        indices = {name: header.index(name) for name in ("NGA", "PolID", "Time")}
        times: set[int] = set()
        polities: set[str] = set()
        ngas: set[str] = set()
        time_series_rows = 0
        for row in time_rows:
            if row[indices["PolID"]] is None or row[indices["Time"]] is None:
                continue
            time_series_rows += 1
            times.add(_as_int(row[indices["Time"]]))
            polities.add(str(row[indices["PolID"]]))
            ngas.add(str(row[indices["NGA"]]))

        population_sheet = workbook["Equinox2020_CanonDat"]
        population_rows = population_sheet.values
        population_header = tuple(next(population_rows))
        population_columns = {
            name: population_header.index(name)
            for name in ("Polity", "Variable", "Value.From", "Date.From")
        }
        polity_population_count = 0
        polity_population_ids: set[str] = set()
        missing_start_dates = 0
        textual_start_values = 0
        for row in population_rows:
            if row[population_columns["Variable"]] != "Polity Population":
                continue
            polity_population_count += 1
            polity_population_ids.add(str(row[population_columns["Polity"]]))
            missing_start_dates += row[population_columns["Date.From"]] is None
            value = row[population_columns["Value.From"]]
            textual_start_values += value is not None and not isinstance(value, (int, float))
        workbook.close()

    return {
        "filename": path.name,
        "sha256": sha256(path),
        "workbook_member": workbook_member,
        "sheets": sheet_profiles,
        "tsdat123": {
            "rows": time_series_rows,
            "polity_ids": len(polities),
            "ngas": len(ngas),
            "time_range": [min(times), max(times)],
            "sentinel_review": "minus_999_and_fractional_codings_must_not_be_treated_as_plain_values",
        },
        "source_support": "seshat_polity_nga_time_not_modern_country_year",
        "polity_population": {
            "rows": polity_population_count,
            "polity_ids": len(polity_population_ids),
            "missing_date_from": missing_start_dates,
            "textual_value_from": textual_start_values,
            "publication_status": "date_value_and_polity_geometry_crosswalk_review_required",
        },
        "publication_status": "variable_selection_and_native_status_adapter_pending",
    }


def build_profile(raw_root: Path, candidate_lock: Path) -> dict[str, Any]:
    lock = json.loads(candidate_lock.read_text(encoding="utf-8"))
    economics = WorldBankEconomicsAdapter(
        gdp_per_capita=raw_root / "world_bank/2026-09-23/gdp-per-capita-constant-2015-usd.json",
        countries=raw_root / "world_bank/2026-09-23/countries.json",
    )
    vdem_release = json.loads(
        (raw_root / "vdem/v16/github-release.json").read_text(encoding="utf-8")
    )
    clio_workbooks = sorted(
        raw_root / item["relative_path"]
        for item in lock["artifacts"]
        if item["source_id"] == "clio_infra" and item["kind"] == "xlsx"
    )
    return {
        "schema_version": "0.1.0",
        "status": "candidate_sources_profiled_not_all_publication_ready",
        "candidate_lock": {
            "filename": candidate_lock.name,
            "sha256": sha256(candidate_lock),
            "artifact_count": len(lock["artifacts"]),
            "blocked_source_count": len(lock["blocked_sources"]),
        },
        "profiles": {
            "world_bank_economics": {
                **economics.profile(),
                "publication_status": "published_first_modern_slice_in_r7",
            },
            "clio_infra": {
                "workbooks": [
                    _clio_profile(path) for path in clio_workbooks
                ],
                "publication_status": "review_only_terms_and_boundary_support_unresolved",
            },
            "vdem": {
                "release": vdem_release.get("tag_name"),
                "published_at": vdem_release.get("published_at"),
                "rdata_sha256": sha256(raw_root / "vdem/v16/vdem.RData"),
                "codebook_sha256": sha256(raw_root / "vdem/v16/codebook_v16.pdf"),
                "first_candidate_variable": {
                    "id": "v2x_suffr",
                    "label": "Share of population with suffrage",
                    "support": "V-Dem country-year",
                    "meaning": "de_jure_share_of_adult_citizens_with_legal_right_to_vote",
                    "scale": [0, 1],
                    "years": [1789, 2025],
                    "codebook_printed_pages": [50, 51],
                    "not_practice_measure": True,
                },
                "publication_status": "large_rdata_parse_and_entity_audit_pending",
            },
            "seshat": _seshat_profile(raw_root / "seshat/equinox-v1/Equinox_Data-v.1.zip"),
            "ucdp": {
                "ged": _ucdp_ged_profile(raw_root / "ucdp/26.1/ged261-csv.zip"),
                "armed_conflict": _ucdp_acd_profile(
                    raw_root / "ucdp/26.1/ucdp-prio-acd-261-csv.zip"
                ),
                "publication_status": "products_kept_separate_adapter_pending",
            },
        },
        "reconciliation_rules": [
            "keep_source_observations_immutable",
            "never_join_historical_and_modern_entities_by_name_alone",
            "keep_statistical_support_political_context_and_display_geometry_separate",
            "do_not_splice_sources_across_time_without_a_versioned_transform",
            "preserve_missing_absent_uncoded_disputed_and_not_applicable_states",
            "source_ranges_and_low_best_high_fields_are_not_statistical_confidence_intervals",
        ],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw-root", type=Path, required=True)
    parser.add_argument("--candidate-lock", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    write_json(args.output, build_profile(args.raw_root, args.candidate_lock))
    print(args.output)


if __name__ == "__main__":
    main()
